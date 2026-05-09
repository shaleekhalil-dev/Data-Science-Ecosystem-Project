from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_ecosystem_report():
    doc = Document()
    
    # إعدادات الخطوط
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Segoe UI'
    font.size = Pt(11)

    # العنوان الرئيسي للمشروع الثالث
    title = doc.add_heading('Data Science Ecosystem: Consumer Behavior Analysis', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. الرؤية الإنسانية
    doc.add_heading('1. Human-Centric Data Vision', level=1)
    doc.add_paragraph(
        "This project explores the intersection of data engineering and psychological capital. "
        "By analyzing user journeys, we aim to build digital environments that are not only "
        "efficient but also humanized."
    )

    # 2. المنهجية التقنية
    doc.add_heading('2. Ecosystem Methodology', level=1)
    doc.add_paragraph(
        "A synthetic big data environment was simulated featuring 500 unique user sessions. "
        "Key metrics include device segmentation, session depth, and conversion probability models."
    )

    # 3. التحليل المرئي
    doc.add_heading('3. Behavioral Insights', level=1)
    
    # الرسم الأول: تحويل الأجهزة
    doc.add_heading('3.1 Device-Based Conversion Analysis', level=2)
    if os.path.exists('figures/device_conversion.png'):
        doc.add_picture('figures/device_conversion.png', width=Inches(5.0))
        doc.add_paragraph("Figure 1: Comparison of conversion rates between Mobile, Desktop, and Tablet users.")

    doc.add_page_break()

    # الرسم الثاني: تشتت السلوك
    doc.add_heading('3.2 Session Engagement vs. Conversion', level=2)
    if os.path.exists('figures/behavior_scatter.png'):
        doc.add_picture('figures/behavior_scatter.png', width=Inches(5.0))
        doc.add_paragraph("Figure 2: Correlation between session duration and pages visited, highlighting successful conversions.")

    # 4. الخلاصة والتوصيات
    doc.add_heading('4. Strategic Conclusions', level=1)
    doc.add_paragraph(
        "The ecosystem demonstrates that digital resilience is built through understanding "
        "user behavior. Optimization should focus on the psychological 'turning points' "
        "identified in the engagement scatter plots."
    )

    output_path = 'outputs/Data_Science_Ecosystem_Report.docx'
    doc.save(output_path)
    print(f"Ecosystem Word document generated at: {output_path}")

if __name__ == "__main__":
    create_ecosystem_report()